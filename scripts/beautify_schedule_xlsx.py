# -*- coding: utf-8 -*-
"""Beautify exported 各平台账号内容排期.xlsx for portfolio use."""

from datetime import datetime, time, date
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

SRC = r'C:\Users\Lenovo\Desktop\各平台账号内容排期.xlsx'
DST = r'C:\Users\Lenovo\Desktop\各平台账号内容排期_美化版.xlsx'

wb = load_workbook(SRC)
ws = wb.active
ws.title = '各平台账号内容排期'

FILL_LEGEND = PatternFill('solid', fgColor='FFF8F1')
FILL_DATE_HDR = PatternFill('solid', fgColor='F0EBE3')
FILL_PUB_HDR = PatternFill('solid', fgColor='E8E2D8')
FILL_XHS = PatternFill('solid', fgColor='9DD4F0')
FILL_BILI = PatternFill('solid', fgColor='98D7B6')
FILL_DY = PatternFill('solid', fgColor='F5D76E')
FILL_WB = PatternFill('solid', fgColor='F5B7C5')
FILL_SUB = PatternFill('solid', fgColor='F7F3EC')
FILL_GOOD = PatternFill('solid', fgColor='FFEEAD')
FILL_BAD = PatternFill('solid', fgColor='FFC9C7')
FILL_SEP = PatternFill('solid', fgColor='FAFAF8')
FILL_WHITE = PatternFill('solid', fgColor='FFFFFF')

TIME_FILLS = {
    '17:30': PatternFill('solid', fgColor='F6C89A'),
    '19:00': PatternFill('solid', fgColor='D5C6EB'),
    '17:00': PatternFill('solid', fgColor='B8D9F0'),
    '18:00': PatternFill('solid', fgColor='F5B7C5'),
    '18:30': PatternFill('solid', fgColor='F5B7C5'),
    '12:00': PatternFill('solid', fgColor='C5E6B8'),
    '14:00': PatternFill('solid', fgColor='F7E08A'),
    '15:00': PatternFill('solid', fgColor='F7E08A'),
    '13:00': PatternFill('solid', fgColor='F7E08A'),
    '10:00': PatternFill('solid', fgColor='C9E4F5'),
    '11:00': PatternFill('solid', fgColor='C9E4F5'),
    '00:00': PatternFill('solid', fgColor='E8D48A'),
    '0:00': PatternFill('solid', fgColor='E8D48A'),
}

TYPE_FILLS = {
    'ID视频': PatternFill('solid', fgColor='C5E6B8'),
    'MV': PatternFill('solid', fgColor='B8D9F0'),
    '自制内容': PatternFill('solid', fgColor='BFDFF0'),
    '新歌上线': PatternFill('solid', fgColor='F5C2D0'),
    '打歌': PatternFill('solid', fgColor='E8D48A'),
    'AI MV': PatternFill('solid', fgColor='D5C6EB'),
}

PLATFORM = {
    '小红书': (2, 6, FILL_XHS),
    'B站': (8, 12, FILL_BILI),
    '抖音': (14, 18, FILL_DY),
    '微博': (20, 22, FILL_WB),
}

thin = Side(style='thin', color='C9C2B8')
medium = Side(style='medium', color='A89F93')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
border_hdr = Border(left=thin, right=thin, top=medium, bottom=medium)

font_legend = Font(name='微软雅黑', size=10, color='5C534A')
font_hdr = Font(name='微软雅黑', size=11, bold=True, color='2F2A26')
font_sub = Font(name='微软雅黑', size=9, bold=True, color='4A433C')
font_body = Font(name='微软雅黑', size=9, color='3D342C')
font_num = Font(name='微软雅黑', size=9, color='3D342C')

align_c = Alignment(horizontal='center', vertical='center', wrap_text=True)
align_l = Alignment(horizontal='left', vertical='center', wrap_text=True)


def fmt_time(v):
    if isinstance(v, datetime):
        return v.strftime('%H:%M')
    if isinstance(v, time):
        return v.strftime('%H:%M')
    if isinstance(v, str) and ':' in v:
        parts = v.split(':')
        return f'{int(parts[0]):02d}:{int(parts[1]):02d}'
    return None


def type_fill(text):
    if not isinstance(text, str):
        return None
    t = text.strip()
    if t in TYPE_FILLS:
        return TYPE_FILLS[t]
    for key, fill in TYPE_FILLS.items():
        if key in t:
            return fill
    return PatternFill('solid', fgColor='EDE8E0')


def is_good_or_bad(cell):
    fill = cell.fill
    if not fill or not fill.patternType:
        return None
    try:
        if fill.fgColor and fill.fgColor.type == 'rgb' and fill.fgColor.rgb:
            rgb = str(fill.fgColor.rgb).upper()
        else:
            return None
    except Exception:
        return None
    if 'FFEEAD' in rgb:
        return 'good'
    if 'FFC9C7' in rgb:
        return 'bad'
    return None


flags = {}
for r in range(7, ws.max_row + 1):
    for c in range(1, ws.max_column + 1):
        tag = is_good_or_bad(ws.cell(r, c))
        if tag:
            flags[(r, c)] = tag

date_merges = [str(m) for m in ws.merged_cells.ranges if str(m).startswith('A')]
for m in list(ws.merged_cells.ranges):
    ws.unmerge_cells(str(m))

ws.merge_cells('A1:V1')
ws['A1'] = '【黄色】为数据较好作品　　【红色】为数据偏低作品'
ws['A1'].font = font_legend
ws['A1'].fill = FILL_LEGEND
ws['A1'].alignment = Alignment(horizontal='left', vertical='center')

for c in range(1, 24):
    ws.cell(2, c).value = None
    ws.cell(3, c).value = None

ws['A4'] = '日期'
ws['A4'].font = font_hdr
ws['A4'].fill = FILL_DATE_HDR
ws['A4'].alignment = align_c
ws['A4'].border = border_hdr

ws.merge_cells('B4:V4')
ws['B4'] = '发布内容'
ws['B4'].font = font_hdr
ws['B4'].fill = FILL_PUB_HDR
ws['B4'].alignment = align_c
ws['B4'].border = border_hdr

for c in range(1, 24):
    if c != 1:
        ws.cell(5, c).value = None

ws['A5'] = None
ws['A5'].fill = FILL_DATE_HDR
ws['A5'].border = border

for name, (c1, c2, fill) in PLATFORM.items():
    if c1 != c2:
        ws.merge_cells(start_row=5, start_column=c1, end_row=5, end_column=c2)
    for c in range(c1, c2 + 1):
        cell = ws.cell(5, c)
        cell.fill = fill
        cell.border = border_hdr
        cell.font = font_hdr
        cell.alignment = align_c
    ws.cell(5, c1).value = name

for c in (7, 13, 19):
    ws.cell(5, c).fill = FILL_SEP
    ws.cell(5, c).border = border

sub_labels = ['发布时间', '内容类型', '内容', '浏览量', '7天内获赞']
for base in (2, 8, 14):
    for i, label in enumerate(sub_labels):
        cell = ws.cell(6, base + i)
        cell.value = label
        cell.font = font_sub
        cell.fill = FILL_SUB
        cell.alignment = align_c
        cell.border = border

for i, label in enumerate(['发布时间', '内容类型', '内容']):
    cell = ws.cell(6, 20 + i)
    cell.value = label
    cell.font = font_sub
    cell.fill = FILL_SUB
    cell.alignment = align_c
    cell.border = border

for c in (7, 13, 19, 23):
    ws.cell(6, c).fill = FILL_SEP
    ws.cell(6, c).border = border

ws['A6'] = None
ws['A6'].fill = FILL_DATE_HDR
ws['A6'].border = border

for r in range(7, ws.max_row + 1):
    dcell = ws.cell(r, 1)
    dcell.font = font_body
    dcell.alignment = align_c
    dcell.border = border
    dcell.fill = FILL_DATE_HDR if dcell.value else FILL_WHITE
    if isinstance(dcell.value, (datetime, date)):
        dcell.number_format = 'YYYY/M/D'

    for c1, c2 in ((2, 6), (8, 12), (14, 18), (20, 22)):
        for c in range(c1, c2 + 1):
            cell = ws.cell(r, c)
            cell.font = font_body
            cell.alignment = align_c
            cell.border = border

            if c in (2, 8, 14, 20):
                tlabel = fmt_time(cell.value)
                if isinstance(cell.value, (datetime, time)):
                    cell.number_format = 'HH:MM'
                cell.fill = TIME_FILLS.get(tlabel, FILL_WHITE) if tlabel else FILL_WHITE
            elif c in (3, 9, 15, 21):
                tf = type_fill(cell.value)
                cell.fill = tf if (cell.value and tf) else FILL_WHITE
            elif c in (4, 10, 16, 22):
                cell.alignment = align_l
                cell.fill = FILL_WHITE
            elif c in (5, 6, 11, 12, 17, 18):
                cell.font = font_num
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0'
                cell.fill = FILL_WHITE

            tag = flags.get((r, c))
            if tag == 'good':
                cell.fill = FILL_GOOD
            elif tag == 'bad':
                cell.fill = FILL_BAD

    for c in (7, 13, 19):
        cell = ws.cell(r, c)
        cell.fill = FILL_SEP
        cell.border = border
        cell.value = None

for m in date_merges:
    try:
        ws.merge_cells(m)
    except Exception:
        pass

ws.column_dimensions['A'].width = 12
for c in range(2, 24):
    letter = get_column_letter(c)
    if c in (7, 13, 19):
        ws.column_dimensions[letter].width = 1.5
    elif c in (4, 10, 16, 22):
        ws.column_dimensions[letter].width = 16
    elif c in (5, 6, 11, 12, 17, 18):
        ws.column_dimensions[letter].width = 11
    elif c in (3, 9, 15, 21):
        ws.column_dimensions[letter].width = 12
    else:
        ws.column_dimensions[letter].width = 10

ws.row_dimensions[1].height = 22
ws.row_dimensions[4].height = 22
ws.row_dimensions[5].height = 24
ws.row_dimensions[6].height = 20
for r in range(7, ws.max_row + 1):
    ws.row_dimensions[r].height = 18

ws.freeze_panes = 'B7'
ws.sheet_view.showGridLines = False
ws.print_title_rows = '1:6'

wb.save(DST)
print('saved', DST)
